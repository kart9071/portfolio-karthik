"""Natural-language expense capture.

POST a sentence like "spent 250 on groceries via UPI yesterday, mom paid 1200
for the electricity bill" and Gemini turns it into structured rows, which land
in SQLite and in the Expenses sheet of an .xlsx laid out exactly like the one
the desktop expense tracker writes.
"""

import json
import os
import sqlite3
from datetime import datetime, date, timedelta

from flask import Blueprint, g, jsonify, request, send_file

from auth import require_login

bp = Blueprint("expenses", __name__)


class UpstreamError(RuntimeError):
    """The Gemini call itself failed - quota, billing, auth or network."""


HERE = os.path.dirname(os.path.abspath(__file__))
DB = os.environ.get("EXPENSES_DB", os.path.join(HERE, "expenses.db"))
XLSX = os.environ.get("EXPENSES_XLSX", os.path.join(HERE, "Expenses.xlsx"))
# flash-lite handles this extraction as accurately as the full models at a
# fraction of the latency and cost. The unversioned alias is deliberate:
# gemini-2.5-flash was retired for new API keys and broke this endpoint, so
# track the current lite model rather than pin one. Set GEMINI_MODEL to pin.
MODEL = os.environ.get("GEMINI_MODEL", "gemini-flash-lite-latest")

# Same controlled vocabularies as expense-tracker/config.json, so rows written
# here stay compatible with the sheet the desktop tool already maintains.
PEOPLE = ["Me", "Mom"]
CATEGORIES = [
    "Groceries", "Vegetables & Fruits", "Milk & Dairy", "Snacks", "Eating Out",
    "Transport", "Fuel", "Medical", "Mobile / Internet", "Electricity & Bills",
    "Rent", "Household Items", "Clothing", "Education", "Gifts",
    "Entertainment", "Other",
]
PAYMENT_MODES = ["Cash", "UPI", "Card", "Bank Transfer", "Other"]

# Column order of the Expenses sheet. Do not reorder - the desktop tool reads
# these positionally.
HEADERS = ["Date", "Person", "Category", "Item / Details", "Amount",
           "Payment Mode", "Notes", "Entered On"]


bp.before_request(require_login)


def get_db():
    conn = sqlite3.connect(DB)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with get_db() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS expenses (
                id           INTEGER PRIMARY KEY AUTOINCREMENT,
                spent_on     TEXT    NOT NULL,
                person       TEXT    NOT NULL,
                category     TEXT    NOT NULL,
                item         TEXT,
                amount       REAL    NOT NULL,
                payment_mode TEXT    NOT NULL,
                notes        TEXT,
                prompt       TEXT    NOT NULL,
                entered_on   TEXT    NOT NULL
            )
        """)
        conn.commit()


# ---------------------------------------------------------------- parsing

RESPONSE_SCHEMA = {
    "type": "OBJECT",
    "properties": {
        "expenses": {
            "type": "ARRAY",
            "items": {
                "type": "OBJECT",
                "properties": {
                    "date": {"type": "STRING", "description": "YYYY-MM-DD"},
                    "person": {"type": "STRING", "enum": PEOPLE},
                    "category": {"type": "STRING", "enum": CATEGORIES},
                    "item": {"type": "STRING"},
                    "amount": {"type": "NUMBER"},
                    "payment_mode": {"type": "STRING", "enum": PAYMENT_MODES},
                    "notes": {"type": "STRING"},
                },
                "required": ["date", "person", "category", "amount", "payment_mode"],
            },
        }
    },
    "required": ["expenses"],
}


def _system_instruction(default_person):
    today = date.today()
    return (
        "You extract expense records from a short note about a household that "
        "tracks two people's spending, Me (Karthik) and Mom.\n"
        f"Today is {today.isoformat()} ({today.strftime('%A')}). "
        f"Yesterday was {(today - timedelta(days=1)).isoformat()}. "
        "Resolve relative dates against those, and assume the current year when "
        "a date gives only day and month. Never return a future date.\n"
        "When one note lists several expenses and only some carry a date, the "
        "undated ones take the last date mentioned before them, or today when "
        "the note gives no date at all.\n"
        f"The note was written by {default_person}, so an expense that names no "
        f"owner belongs to Person='{default_person}'. An explicit mention of "
        "mother/mom/amma still means Person='Mom', and 'I' or 'me' means the writer.\n"
        "Amounts are Indian rupees; return the number only, no symbol.\n"
        "Pick the closest category from the allowed list, 'Other' if none fit. "
        "Default Payment Mode to 'UPI' when unstated.\n"
        "Put the specific thing bought in 'item' and anything else worth keeping "
        "in 'notes'. Return one entry per distinct expense."
    )


def parse_expenses(prompt, default_person="Me"):
    """Ask Gemini to turn free text into expense rows. Returns a list of dicts."""
    # Imported lazily and the key read per-request, so the service still boots
    # (and CI still passes) on a host with no GEMINI_API_KEY set.
    from google import genai
    from google.genai import types

    # The Google SDK honours either name, so accept both rather than making a
    # host that already has GOOGLE_API_KEY set define a second variable.
    api_key = os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY")
    if not api_key:
        raise RuntimeError("Set GEMINI_API_KEY (or GOOGLE_API_KEY) on this host")

    client = genai.Client(api_key=api_key)
    try:
        resp = client.models.generate_content(
            model=MODEL,
            contents=prompt,
            config=types.GenerateContentConfig(
                system_instruction=_system_instruction(default_person),
                response_mime_type="application/json",
                response_schema=RESPONSE_SCHEMA,
                temperature=0,
            ),
        )
    except Exception as e:
        # Quota, billing and auth failures are not the prompt's fault - keep
        # them distinguishable so the caller is not sent to debug their wording.
        raise UpstreamError(str(e))

    payload = json.loads(resp.text)
    return payload.get("expenses") or []


def clean_row(raw):
    """Coerce one model-produced dict into the shape the sheet and DB expect."""
    try:
        amount = float(raw.get("amount"))
    except (TypeError, ValueError):
        raise ValueError(f"amount is not a number: {raw.get('amount')!r}")
    if amount <= 0:
        raise ValueError(f"amount must be positive, got {amount}")

    spent_on = (raw.get("date") or "").strip()
    try:
        datetime.strptime(spent_on, "%Y-%m-%d")
    except ValueError:
        raise ValueError(f"date is not YYYY-MM-DD: {spent_on!r}")

    person = (raw.get("person") or "").strip()
    if person not in PEOPLE:
        raise ValueError(f"person must be one of {PEOPLE}, got {person!r}")

    category = (raw.get("category") or "Other").strip()
    if category not in CATEGORIES:
        category = "Other"

    mode = (raw.get("payment_mode") or "UPI").strip()
    if mode not in PAYMENT_MODES:
        mode = "Other"

    return {
        "spent_on": spent_on,
        "person": person,
        "category": category,
        "item": (raw.get("item") or "").strip() or None,
        "amount": amount,
        "payment_mode": mode,
        "notes": (raw.get("notes") or "").strip() or None,
    }


# ---------------------------------------------------------------- storage

def append_to_xlsx(rows):
    """Append rows to the Expenses sheet, creating the workbook if absent."""
    import openpyxl

    if os.path.exists(XLSX):
        wb = openpyxl.load_workbook(XLSX)
        ws = wb["Expenses"] if "Expenses" in wb.sheetnames else wb.create_sheet("Expenses")
        if ws.max_row == 1 and all(c.value is None for c in ws[1]):
            ws.append(HEADERS)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Expenses"
        ws.append(HEADERS)

    for r in rows:
        ws.append([
            datetime.strptime(r["spent_on"], "%Y-%m-%d"),
            r["person"],
            r["category"],
            r["item"],
            r["amount"],
            r["payment_mode"],
            r["notes"],
            datetime.strptime(r["entered_on"], "%Y-%m-%d %H:%M:%S"),
        ])

    for cell in ws["A"][1:]:
        cell.number_format = "DD-MM-YYYY"
    for cell in ws["H"][1:]:
        cell.number_format = "DD-MM-YYYY HH:MM:SS"

    # Write via a temp file so a crash mid-save cannot truncate the sheet.
    tmp = XLSX + ".tmp"
    wb.save(tmp)
    os.replace(tmp, XLSX)


def rebuild_xlsx():
    """Rewrite the whole Expenses sheet from the database.

    Used after a delete: openpyxl can remove a row, but matching a DB row to a
    sheet row means guessing, and rewriting from the record keeps the two
    genuinely in step. Any other sheets in the workbook are preserved.
    """
    import openpyxl

    with get_db() as conn:
        rows = [dict(r) for r in conn.execute(
            "SELECT * FROM expenses ORDER BY spent_on, id").fetchall()]

    if os.path.exists(XLSX):
        wb = openpyxl.load_workbook(XLSX)
        if "Expenses" in wb.sheetnames:
            del wb["Expenses"]
        ws = wb.create_sheet("Expenses", 0)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Expenses"

    ws.append(HEADERS)
    for r in rows:
        ws.append([
            datetime.strptime(r["spent_on"], "%Y-%m-%d"),
            r["person"], r["category"], r["item"], r["amount"],
            r["payment_mode"], r["notes"],
            datetime.strptime(r["entered_on"], "%Y-%m-%d %H:%M:%S"),
        ])
    for cell in ws["A"][1:]:
        cell.number_format = "DD-MM-YYYY"
    for cell in ws["H"][1:]:
        cell.number_format = "DD-MM-YYYY HH:MM:SS"

    tmp = XLSX + ".tmp"
    wb.save(tmp)
    os.replace(tmp, XLSX)


def insert_rows(rows, prompt):
    entered_on = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    saved = []
    with get_db() as conn:
        for r in rows:
            r["entered_on"] = entered_on
            cur = conn.execute(
                """INSERT INTO expenses
                   (spent_on, person, category, item, amount, payment_mode, notes, prompt, entered_on)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (r["spent_on"], r["person"], r["category"], r["item"], r["amount"],
                 r["payment_mode"], r["notes"], prompt, entered_on),
            )
            saved.append({"id": cur.lastrowid, **r})
        conn.commit()
    return saved


# ---------------------------------------------------------------- routes

@bp.route("/api/expense", methods=["POST"])
def add_expense():
    data = request.get_json(silent=True)
    if not data:
        return jsonify({"error": "Invalid JSON"}), 400

    prompt = (data.get("prompt") or "").strip()
    if not prompt:
        return jsonify({"error": "prompt is required"}), 422

    try:
        raw_rows = parse_expenses(prompt, g.user["person"])
    except UpstreamError as e:
        return jsonify({"error": "Gemini rejected the request", "detail": str(e)}), 502
    except RuntimeError as e:
        return jsonify({"error": str(e)}), 503
    except Exception as e:
        return jsonify({"error": f"could not read the model's reply: {e}"}), 502

    if not raw_rows:
        return jsonify({"error": "no expense found in that prompt", "prompt": prompt}), 422

    try:
        rows = [clean_row(r) for r in raw_rows]
    except ValueError as e:
        return jsonify({"error": f"model returned an unusable row: {e}"}), 502

    saved = insert_rows(rows, prompt)

    # The sheet is a convenience copy; the DB is the record. If the write fails
    # say so rather than losing the rows that already committed.
    xlsx_ok, xlsx_error = True, None
    try:
        append_to_xlsx(saved)
    except Exception as e:
        xlsx_ok, xlsx_error = False, str(e)

    total = round(sum(r["amount"] for r in saved), 2)
    for r in saved:
        print(f"[EXPENSE] {r['spent_on']} {r['person']:4} {r['category']:20} {r['amount']:>9.2f}")

    body = {"success": True, "count": len(saved), "total_added": total,
            "expenses": saved, "xlsx_updated": xlsx_ok}
    if xlsx_error:
        body["xlsx_error"] = xlsx_error
    return jsonify(body), 201


@bp.route("/api/expenses", methods=["GET"])
def list_expenses():
    person = (request.args.get("person") or "").strip()
    date_from = (request.args.get("from") or "").strip()
    date_to = (request.args.get("to") or "").strip()
    month = (request.args.get("month") or "").strip()  # YYYY-MM
    on_date = (request.args.get("date") or "").strip()  # YYYY-MM-DD

    sql = "SELECT * FROM expenses WHERE 1=1"
    params = []
    if person:
        sql += " AND person = ? COLLATE NOCASE"
        params.append(person)
    if on_date:
        sql += " AND spent_on = ?"
        params.append(on_date)
    if month:
        sql += " AND substr(spent_on, 1, 7) = ?"
        params.append(month)
    if date_from:
        sql += " AND spent_on >= ?"
        params.append(date_from)
    if date_to:
        sql += " AND spent_on <= ?"
        params.append(date_to)
    sql += " ORDER BY spent_on DESC, id DESC"

    with get_db() as conn:
        rows = [dict(r) for r in conn.execute(sql, params).fetchall()]

    by_person = {}
    by_category = {}
    for r in rows:
        by_person[r["person"]] = round(by_person.get(r["person"], 0) + r["amount"], 2)
        by_category[r["category"]] = round(by_category.get(r["category"], 0) + r["amount"], 2)

    return jsonify({
        "count": len(rows),
        "total": round(sum(r["amount"] for r in rows), 2),
        "by_person": by_person,
        "by_category": by_category,
        "filters": {"person": person or None, "month": month or None,
                    "date": on_date or None,
                    "from": date_from or None, "to": date_to or None},
        "expenses": rows,
    })


@bp.route("/api/expenses/<int:expense_id>", methods=["DELETE"])
def delete_expense(expense_id):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM expenses WHERE id = ?", (expense_id,)).fetchone()
        if not row:
            return jsonify({"error": "no expense with that id"}), 404
        deleted = dict(row)
        conn.execute("DELETE FROM expenses WHERE id = ?", (expense_id,))
        conn.commit()

    xlsx_ok, xlsx_error = True, None
    try:
        rebuild_xlsx()
    except Exception as e:
        xlsx_ok, xlsx_error = False, str(e)

    print(f"[DELETED] {deleted['spent_on']} {deleted['person']} "
          f"{deleted['category']} {deleted['amount']} by {g.user['username']}")

    body = {"success": True, "deleted": deleted, "xlsx_updated": xlsx_ok}
    if xlsx_error:
        body["xlsx_error"] = xlsx_error
    return jsonify(body)


@bp.route("/api/expenses/export", methods=["GET"])
def export_expenses():
    if not os.path.exists(XLSX):
        return jsonify({"error": "no spreadsheet yet - add an expense first"}), 404
    return send_file(XLSX, as_attachment=True, download_name="Expenses.xlsx")


init_db()
